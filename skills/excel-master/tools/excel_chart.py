#!/usr/bin/env python3
"""excel_chart.py — 在已有 xlsx 上加图表(柱/折线/面积/双轴)。

Spec 格式:
{
  "sheet": "趋势数据",                    # 目标 sheet 名(必须已存在)
  "charts": [
    {
      "type": "combo_dual_axis_line",     # line | bar | column | area | combo_dual_axis_line
      "anchor": "H2",                     # 图表锚点单元格
      "title": "交易额 vs 出口占比",
      "primary": {                        # 双轴图的主系列
        "series_col": "B",
        "header_row": 1,
        "data_row_start": 2,
        "data_row_end": 11,
        "name": "交易额"
      },
      "secondary": {                      # 双轴图的副系列
        "series_col": "E",
        "header_row": 1,
        "data_row_start": 2,
        "data_row_end": 11,
        "name": "出口占比"
      },
      "categories_col": "A",
      "categories_row_start": 2,
      "categories_row_end": 11
    }
  ]
}

单系列图(line/bar/column/area)只填 series + categories。
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
    from openpyxl.chart import AreaChart, BarChart, LineChart, Reference
    from openpyxl.chart.layout import Layout, ManualLayout
except ImportError:
    print("openpyxl not installed", file=sys.stderr)
    sys.exit(2)


PRIMARY_COLOR = "4472C4"   # 蓝
SECONDARY_COLOR = "ED7D31"  # 橙
AREA_COLOR = "8FAADC"


def make_single_chart(ws, ctype: str, c: dict[str, Any]):
    if ctype == "line":
        chart = LineChart()
    elif ctype == "bar":
        chart = BarChart()
        chart.type = "bar"
    elif ctype == "column":
        chart = BarChart()
        chart.type = "col"
    elif ctype == "area":
        chart = AreaChart()
    else:
        raise ValueError(f"unknown chart type: {ctype}")

    col = c["series_col"]
    col_idx = ord(col.upper()) - ord("A") + 1
    data_ref = Reference(
        ws,
        min_col=col_idx,
        max_col=col_idx,
        min_row=c["header_row"],
        max_row=c["data_row_end"],
    )
    chart.add_data(data_ref, titles_from_data=True)

    cats_ref = Reference(
        ws,
        min_col=ord(c["categories_col"].upper()) - ord("A") + 1,
        max_col=ord(c["categories_col"].upper()) - ord("A") + 1,
        min_row=c["categories_row_start"],
        max_row=c["categories_row_end"],
    )
    chart.set_categories(cats_ref)

    chart.title = c.get("title", "")
    chart.width = c.get("width", 18)
    chart.height = c.get("height", 10)
    return chart


def make_dual_axis(ws, c: dict[str, Any]) -> Any:
    """双轴折线:主+副两个 LineChart 合并,副轴交叉到右边。"""
    primary_spec = c["primary"]
    secondary_spec = c["secondary"]

    cat_col = ord(c["categories_col"].upper()) - ord("A") + 1

    # 主图
    chart1 = LineChart()
    p_idx = ord(primary_spec["series_col"].upper()) - ord("A") + 1
    p_ref = Reference(
        ws,
        min_col=p_idx,
        max_col=p_idx,
        min_row=primary_spec["header_row"],
        max_row=primary_spec["data_row_end"],
    )
    chart1.add_data(p_ref, titles_from_data=True)
    cats_ref = Reference(
        ws,
        min_col=cat_col,
        max_col=cat_col,
        min_row=c["categories_row_start"],
        max_row=c["categories_row_end"],
    )
    chart1.set_categories(cats_ref)

    # 副图
    chart2 = LineChart()
    s_idx = ord(secondary_spec["series_col"].upper()) - ord("A") + 1
    s_ref = Reference(
        ws,
        min_col=s_idx,
        max_col=s_idx,
        min_row=secondary_spec["header_row"],
        max_row=secondary_spec["data_row_end"],
    )
    chart2.add_data(s_ref, titles_from_data=True)
    chart2.y_axis.axId = 200
    chart2.y_axis.crosses = "max"
    chart2.y_axis.title = secondary_spec.get("name", "")

    chart1 += chart2  # 合并

    chart1.title = c.get("title", "")
    chart1.y_axis.title = primary_spec.get("name", "")
    chart1.width = c.get("width", 20)
    chart1.height = c.get("height", 10)
    return chart1


def main() -> int:
    p = argparse.ArgumentParser(description="Add charts to an existing xlsx")
    p.add_argument("--xlsx", required=True, help="path to existing .xlsx")
    p.add_argument("--spec", required=True, help="path to chart spec JSON")
    args = p.parse_args()

    xlsx_path = Path(args.xlsx)
    spec_path = Path(args.spec)
    if not xlsx_path.exists():
        print(f"xlsx not found: {xlsx_path}", file=sys.stderr)
        return 1
    if not spec_path.exists():
        print(f"spec not found: {spec_path}", file=sys.stderr)
        return 1

    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    sheet_name = spec["sheet"]
    charts_spec = spec.get("charts", [])

    wb = load_workbook(str(xlsx_path))
    if sheet_name not in wb.sheetnames:
        print(f"sheet '{sheet_name}' not in workbook (have: {wb.sheetnames})", file=sys.stderr)
        return 1
    ws = wb[sheet_name]

    added = []
    for c in charts_spec:
        ctype = c["type"]
        anchor = c.get("anchor", "H2")
        if ctype == "combo_dual_axis_line":
            chart = make_dual_axis(ws, c)
        else:
            chart = make_single_chart(ws, ctype, c)
        ws.add_chart(chart, anchor)
        added.append({"type": ctype, "anchor": anchor, "title": c.get("title", "")})

    wb.save(str(xlsx_path))

    json.dump(
        {"xlsx": str(xlsx_path.resolve()), "sheets": len(wb.sheetnames), "charts_added": added},
        sys.stdout,
        ensure_ascii=False,
    )
    print()
    return 0


if __name__ == "__main__":
    sys.exit(main())